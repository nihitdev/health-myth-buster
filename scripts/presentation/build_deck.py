"""Fill the user's official template without changing its first slide.
Outputs remain local; the original template and downloaded narration are not published.
"""
import base64, html, json, math, re, sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.dml import MSO_FILL

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'deliverables/health-myth-buster'
TEMPLATE=Path(sys.argv[1]) if len(sys.argv)>1 else Path.home()/'Downloads/PPT_Template_Grade6_AI_HealthMythBuster.pptx'
OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(TEMPLATE)
assert len(prs.slides)==9,'Inspect a changed template before filling it.'
INK='1F2937'; BLUE='005A9E'; TEAL='087B70'; MUTED='526970'; AMBER='9B5712'
PUBLIC='https://health-myth-buster-amber.vercel.app/'

def fill(shape,text,size=12,color=INK,bold=False):
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(.04);tf.margin_right=Inches(.04)
    tf.margin_top=Inches(.02);tf.margin_bottom=Inches(.02)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(4);p.line_spacing=1.08
        r=p.add_run();r.text=line;r.font.name='Calibri';r.font.size=Pt(size)
        r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    return shape

def text(slide,x,y,w,h,value,size=12,color=INK,bold=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    return fill(sh,value,size,color,bold)

def card(slide,x,y,w,h,title,body):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('EDF6F4');sh.line.color.rgb=RGBColor.from_string('CFE1DC')
    text(slide,x+.15,y+.13,w-.3,.35,title,13,TEAL,True)
    text(slide,x+.15,y+.6,w-.3,h-.7,body,12)

def notes(slide,value):slide.notes_slide.notes_text_frame.text=value

def image_fit(slide,file,x,y,w,h):
    iw,ih=Image.open(file).size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
    slide.shapes.add_picture(str(file),Inches(x+(w-nw)/2),Inches(y),width=Inches(nw),height=Inches(nh))

# Slide 1 is deliberately never edited.
s=prs.slides[1]
fill(s.shapes[8],'team_1127_0032',14,TEAL,True)
fill(s.shapes[10],'Birla Open Minds International School\nIndore, India',11)
fill(s.shapes[12],'Nihit Sunhare — built and developed the tool\nRudransh Sahu — team member\nPriyushi Agrawal — team member\nAhika Porwal — team member',10.7)
s.shapes[12].text_frame.paragraphs[0].runs[0].font.bold=True
fill(s.shapes[14],PUBLIC,10,TEAL)
s.shapes[14].text_frame.paragraphs[0].runs[0].hyperlink.address=PUBLIC
fill(s.shapes[16],'React + TypeScript + Vite\nLocal database; transparent keyword rules',10.5)
fill(s.shapes[17],'Reviewed entries / pending candidates',9.5,BLUE,True)
fill(s.shapes[18],'36 reviewed\n31 pending',10,TEAL,True)
notes(s,'Team name, school and website were supplied by Nihit. Development credit is based on his stated contribution. No other student development roles are invented. The public URL was checked in an isolated signed-out Chromium browser: HTTP 200, FALSE and UNVERIFIED examples, /library and /presentation direct routes. Still check on another physical phone before submission. The repository is https://github.com/nihitdev/health-myth-buster .')

s=prs.slides[2]
card(s,.4,1.98,4.5,2.47,'THE PROBLEM','Health forwards can sound convincing without reliable evidence.\nThey spread through chats and social feeds.\nBelieving them can delay reliable information or medical help.')
card(s,5.1,1.98,4.5,2.47,'OUR USERS & RESPONSE','Students, families, teachers and NGO volunteers.\nOne claim → one clear verdict, a short explanation and official sources.\nUnknown or unclear wording stays UNVERIFIED.')
text(s,.45,4.56,9.1,.55,'Database example: “Antibiotics cure the common cold.” → FALSE\nBefore submission: add one real local forward collected with permission; remove personal details.',10,AMBER)
notes(s,'Do not say that this example was received by a family member unless that actually happened. The official task asks for real collected messages. Add a genuine anonymized example with permission; this field remains unfinished until that evidence is collected. No names or phone numbers from a forward should be saved.')

s=prs.slides[3]
for idx,x,filename,caption in [
    (8,.4,'screen-1-input.png','Enter or paste one health claim.\nExamples also work with one click.'),
    (10,3.5,'screen-2-false.png','A real FALSE verdict, its reason\nand two official source links.'),
    (12,6.6,'screen-3-unverified.png','An unknown claim stays UNVERIFIED.\nThe tool does not guess.')]:
    sh=s.shapes[idx];sh.left=Inches(x);sh.top=Inches(4.62);sh.width=Inches(2.93);sh.height=Inches(.5)
    fill(sh,caption,10.3)
    image_fit(s,OUT/'screens'/filename,x,2.31,2.93,2.2)
notes(s,'All three images are actual screenshots of the current app using its real engine, not edited verdicts or invented evaluation results. Full-resolution originals are included in the screens folder. Slide 2 contains the public tool link.')

s=prs.slides[4]
fill(s.shapes[8],'Enter one English health claim, up to 300 characters.\nOr tap a ready-made example.\nExample: “Do antibiotics work for the flu?”',12)
fill(s.shapes[11],'Normalize case, punctuation, spelling and word forms.\nCheck reviewed wording patterns or all required keyword groups.\nUse an 85% coverage threshold; extra unsupported content vetoes a match.\nUnreviewed negation or ambiguity → UNVERIFIED.',11.2)
fill(s.shapes[14],'Show TRUE, FALSE or UNVERIFIED.\nMatched claims include a reason and 2+ official source links.\nUnknown wording gets clarification where possible.\nEvery answer ends:\n“Always ask a doctor.”',12)
notes(s,'The checker is a deterministic rule-based prototype, not machine learning, an LLM or an autonomous medical expert. A rule score is not medical certainty. Opposite-polarity wording is accepted only when explicitly reviewed with its own verdict. All required concepts must match. The numeric threshold alone is not enough: unsupported extra content also vetoes a match. No diagnosis or treatment is provided.')

samples=[
 ('Antibiotics cure the common cold.','antibiotics + cure/treat + cold','FALSE',[
 ('CDC: Antibiotic use','https://www.cdc.gov/antibiotic-use/about/index.html'),('NHS: Antibiotics','https://www.nhs.uk/medicines/antibiotics/')]),
 ('Vaccines cause autism.','vaccine/vaccination + cause + autism','FALSE',[
 ('WHO: Vaccine safety','https://www.who.int/news-room/questions-and-answers/item/vaccines-and-immunization-vaccine-safety'),('NHS: Why vaccination matters','https://www.nhs.uk/vaccinations/why-vaccination-is-important-and-the-safest-way-to-protect-yourself/')]),
 ('HIV spreads through hugging.','HIV + hug/hugging + spread/catch','FALSE',[
 ('WHO: HIV and AIDS','https://www.who.int/news-room/fact-sheets/detail/hiv-aids'),('ANACS: HIV basics (India)','https://anacs.andamannicobar.gov.in/BasicFAQS.php')]),
 ('Exercise lowers heart disease risk.','exercise + reduce/lower + heart disease + risk','TRUE',[
 ('WHO: Physical activity','https://www.who.int/news-room/fact-sheets/detail/physical-activity'),('CDC: Activity benefits','https://www.cdc.gov/physical-activity-basics/benefits/index.html')]),
 ('Boiling water removes toxic chemicals.','boiling + water + chemicals + removes','FALSE',[
 ('CDC: Household water treatment','https://www.cdc.gov/global-water-sanitation-hygiene/about/about-household-water-treatment.html'),('CDC: Water emergencies','https://www.cdc.gov/water-emergency/about/')])]
s=prs.slides[5]
tbl=s.shapes[7].table
tbl.cell(6,0).text='Database size & review'
tbl.cell(6,1).text='36 reviewed entries; 31 pending candidates excluded. Each published verdict has 2+ official sources. Adult/clinical review pending.'
for c,w in zip(tbl.columns,[2.65,2.15,.8,3.6]):c.width=Inches(w)
for j,t in enumerate(['Claim in everyday words','Required keyword concepts','Verdict','Official evidence (click each source)']):tbl.cell(0,j).text=t
for i,(claim,keywords,verdict,refs) in enumerate(samples,1):
    tbl.cell(i,0).text=claim;tbl.cell(i,1).text=keywords;tbl.cell(i,2).text=verdict
    cell=tbl.cell(i,3);cell.text_frame.clear()
    for k,(label,url) in enumerate(refs):
        p=cell.text_frame.paragraphs[0] if k==0 else cell.text_frame.add_paragraph();r=p.add_run();r.text=label;r.hyperlink.address=url
for i,row in enumerate(tbl.rows):
    row.height=Inches(.43 if i<6 else .38)
    for j,cell in enumerate(row.cells):
        cell.margin_left=Inches(.07);cell.margin_right=Inches(.07);cell.margin_top=Inches(.035);cell.margin_bottom=Inches(.02)
        cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('E7F2F8' if i==0 else ('F4F8FA' if i%2 else 'FFFFFF'))
        for p in cell.text_frame.paragraphs:
            p.space_after=Pt(1);p.line_spacing=1.0
            for r in p.runs:r.font.name='Calibri';r.font.size=Pt(9.3);r.font.bold=(i==0 or j==2);r.font.color.rgb=RGBColor.from_string(BLUE if i==0 else (TEAL if j==3 else INK))

notes(s,'Five samples from the real reviewed database. Exact source URLs:\n\n'+'\n\n'.join(claim+'\n'+'\n'.join(label+': '+url for label,url in refs) for claim,_,_,refs in samples)+'\n\nSource-content review was performed during development on 2026-10-22, not independently certified by a clinician. Review notes and unresolved candidates are in SOURCE_AUDIT.md. Some pairs are distinct documents from the same official agency. No claims are promoted merely to reach 50.')

s=prs.slides[6]
fill(s.shapes[9],'Human evaluation: NOT YET MEASURED\nRequired: 5+ other people, 20+ original claims, including unknown claims.\nCorrect / total: record genuine results.\nAccuracy = correct ÷ total × 100.\nDo not substitute software-test counts for user accuracy.',11.6)
fill(s.shapes[10],'Implemented improvements:\n• Reviewed library: 12 → 36 entries.\n• Everyday questions and word forms.\n• Explicitly reviewed negative phrasing.\n• Clarification for broad claims.\nSoftware checks: 247 unit + 8 browser tests passed.\nNext: genuine before/after evaluation.',11.4)
notes(s,'This slide deliberately has no fabricated participant count or accuracy. Collect consented, anonymized evaluations in /testing. Export records, fix failures and rerun identical baseline wordings. Software regression tests are not independent medical validation. The change from 12 to 36 reviewed entries and matching improvements are actual recorded implementation work, not measured user-test outcomes.')

s=prs.slides[7]
fill(s.shapes[8],'School lessons, family awareness and NGO workshops.\nShare the public link or QR code. No login is needed.',11.6)
fill(s.shapes[10],'Free static app; no paid API.\nProposed maintainer: Nihit.\nNew claims need 2+ official sources and adult/health-expert review.',11.3)
fill(s.shapes[12],'Pause before you forward a health tip.\nCheck clear verdicts and read the evidence.\nUse it for awareness—and always ask a doctor.',17,TEAL,True)
notes(s,'Public website: '+PUBLIC+'\nPublic GitHub repository: https://github.com/nihitdev/health-myth-buster\nThe signed-out browser check passed for the homepage, reviewed and unknown claims, and direct Library/Presentation routes. Another physical phone check is still required. Propose updating the reviewed database using the documented review procedure; do not claim a clinician already approved it. The application itself needs no AI API or backend.')

s=prs.slides[8]
for name,idx,bodyidx in [('Nihit Sunhare',6,7),('Rudransh Sahu',9,10),('Priyushi Agrawal',12,13),('Ahika Porwal',15,16)]:
    fill(s.shapes[idx],name,11,BLUE,True)
    fill(s.shapes[bodyidx],'WRITE YOUR OWN 3–4 SENTENCES:\nWhat problem do we solve? How does the tool work? What did I actually do or learn?',10.8,AMBER)
notes(s,'The actual template includes this ninth slide, although the student instruction summary mentions eight. Do not delete it. Every member must independently write their own answer. These are prompts, not generated answers attributed to students. Do not invent personal contributions.')

file=OUT/'Health_Myth_Buster_Official_Template_DRAFT.pptx';prs.save(file)
# Restore the untouched title part byte-for-byte, including its XML declaration.
with ZipFile(TEMPLATE) as original: title_xml=original.read('ppt/slides/slide1.xml')
with ZipFile(file) as built: parts=[(info,built.read(info.filename)) for info in built.infolist()]
tmp=file.with_suffix('.tmp')
with ZipFile(tmp,'w',ZIP_DEFLATED) as archive:
    for info,data in parts: archive.writestr(info,title_xml if info.filename=='ppt/slides/slide1.xml' else data)
tmp.replace(file)
with ZipFile(file) as built:
    assert built.read('ppt/slides/slide1.xml')==title_xml,'Official title slide changed'
assert len(Presentation(file).slides)==9

# Generate a browser-based layout preview from the PowerPoint shapes.
# It is a preview, not a claim to have run Microsoft PowerPoint's renderer.
def color(value,default='FFFFFF'):
    try:
        if value.type==1:return str(value.rgb)
        if value.type==2 and ('BACKGROUND_1' in str(value.theme_color) or 'LIGHT_1' in str(value.theme_color)):return 'FFFFFF'
        return default
    except (AttributeError,TypeError):return default

def tx(tf):
    parts=[]
    for p in tf.paragraphs:
        runs=[]
        for r in p.runs:
            f=r.font;size=f.size.pt if f.size else 11
            style=f'font-size:{size*128/72:.2f}px;color:#{color(f.color,INK)};font-weight:{700 if f.bold else 400}'
            value=html.escape(r.text).replace('\n','<br/>')
            if r.hyperlink.address:value=f'<a href="{html.escape(r.hyperlink.address)}">{value}</a>'
            runs.append(f'<span style="{style}">{value}</span>')
        parts.append('<p>'+''.join(runs)+'</p>')
    return ''.join(parts)

sections=[]
for i,slide in enumerate(prs.slides):
    elements=[]
    bg=color(slide.background.fill.fore_color,'FFFFFF') if slide.background.fill.type==1 else 'FFFFFF'
    for sh in slide.shapes:
        box='position:absolute;left:%.2fpx;top:%.2fpx;width:%.2fpx;height:%.2fpx;' % tuple(getattr(sh,p)/914400*128 for p in ['left','top','width','height'])
        if sh.has_table:
            cells=[]
            for row in sh.table.rows:
                cells.append('<tr>'+''.join(f'<td colspan="{cell.span_width if cell.is_merge_origin else 1}" style="width:{col.width/914400*128}px;background:#{color(cell.fill.fore_color) if cell.fill.type==1 else 'FFFFFF'}">{tx(cell.text_frame)}</td>' for col,cell in zip(sh.table.columns,row.cells) if not cell.is_spanned)+'</tr>')
            elements.append(f'<table style="{box}"> '+''.join(cells)+'</table>');continue
        if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
            b64=base64.b64encode(sh.image.blob).decode();elements.append(f'<img style="{box}" src="data:{sh.image.content_type};base64,{b64}"/>');continue
        fillstyle=''
        if sh.fill.type==MSO_FILL.SOLID:fillstyle=f'background:#{color(sh.fill.fore_color)};'
        if sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:fillstyle+='border-radius:10px;'
        value=tx(sh.text_frame) if sh.has_text_frame else ''
        padding=''
        if sh.has_text_frame:
            tf=sh.text_frame;padding='padding:%.2fpx %.2fpx %.2fpx %.2fpx;' % tuple(getattr(tf,p)/914400*128 for p in ['margin_top','margin_right','margin_bottom','margin_left'])
        elements.append(f'<div style="{box}{fillstyle}{padding}">{value}</div>')
    sections.append(f'<section class="slide" id="slide-{i+1}" style="background:#{bg}">'+''.join(elements)+'</section>')
htmlfile='''<!doctype html><html><head><meta charset="utf-8"><title>Health Myth Buster — official template preview</title><style>
*{box-sizing:border-box}body{margin:0;background:#dce6e3;font-family:Calibri,Arial,sans-serif}.slide{position:relative;width:1280px;height:720px;overflow:hidden;break-after:page;margin:0 auto 24px}p{margin:0 0 6px;line-height:1.12}a{color:inherit;text-decoration:underline}table{border-collapse:collapse;table-layout:fixed}td{border:1px solid #cad9df;padding:5px 8px;vertical-align:top}td p{margin-bottom:2px}@page{size:1280px 720px;margin:0}@media print{body{background:white}.slide{margin:0;break-after:page}}
</style></head><body>'''+''.join(sections)+'</body></html>'
(OUT/'presentation-preview.html').write_text(htmlfile)
(OUT/'deck-manifest.json').write_text(json.dumps({'slides':9,'titleSlideUnchanged':True,'team':'team_1127_0032','school':'Birla Open Minds International School','city':'Indore, India','publicUrl':PUBLIC,'draftReasons':['Genuine participant evaluation not supplied','Real local forward not supplied','Each student must independently complete Slide 9','Adult/clinical evidence review pending'],'templateSource':TEMPLATE.name},indent=2))
print('Saved editable official-template draft:',file)
print('Verified 9 slides; official title-slide XML unchanged.')

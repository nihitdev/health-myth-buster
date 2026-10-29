"""Update presentation with explicitly supplied aggregate counts; never fabricate raw participants."""
import json,base64,html,argparse
from datetime import date
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'deliverables/ppt-update';OUT.mkdir(parents=True,exist_ok=True)
BASE=ROOT/'deliverables/health-myth-buster/Health_Myth_Buster_Official_Template_DRAFT.pptx'
benchmark=json.loads((OUT/'Automated_Benchmark.json').read_text())
parser=argparse.ArgumentParser();parser.add_argument('--correct',type=int,required=True);parser.add_argument('--total',type=int,required=True);parser.add_argument('--participants',type=int,required=True);args=parser.parse_args()
assert args.total>0 and 0<=args.correct<=args.total and args.participants>0
summary={'participants':args.participants,'total':args.total,'correct':args.correct,'incorrect':args.total-args.correct,'accuracyPercent':args.correct/args.total*100,'basis':'Aggregate counts explicitly supplied by the team on '+str(date.today()),'claimLevelRecordsProvided':False,'beforeAfterUserResultsProvided':False}
assert 0<=summary['correct']<=summary['total']
(OUT/'User_Test_Summary.json').write_text(json.dumps(summary,indent=2))
prs=Presentation(BASE)
def fill(shape,value,size=11.5,color='1F2937'):
 tf=shape.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=Inches(.04);tf.margin_right=Inches(.04);tf.margin_top=Inches(.02);tf.margin_bottom=Inches(.02)
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(4);p.line_spacing=1.05
  r=p.add_run();r.text=line;r.font.name='Calibri';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
 return shape
s=prs.slides[2]
for sh in s.shapes:
 if sh.has_text_frame and 'Database example:' in sh.text:
  fill(sh,'Real development input: “Antibiotics cure common illness.” → UNVERIFIED because the illness is unspecified.\nObserved in our testing discussion; this was typed into the tool, not collected as a WhatsApp forward.',9.8,'087B70')
s.notes_slide.notes_text_frame.text='This example is a real claim Nihit entered during development and copied from the Testing Lab into the project conversation. He asked why it was UNVERIFIED. The actual engine leaves the broad statement unverified and now gives clearer wording guidance. It is not presented as a collected family forward. No permission to reproduce a private chat, name or phone number is implied. The statement itself contains no personal information.'
s=prs.slides[6]
fill(s.shapes[9],f'User evaluation — team-reported\n{summary["participants"]} other people tried {summary["total"]} claims.\n{summary["correct"]} correct; {summary["incorrect"]} incorrect.\nAccuracy = {summary["correct"]} ÷ {summary["total"]} × 100 = {summary["accuracyPercent"]:g}%.\nKeep the original claim-by-claim record for review.',12)
fill(s.shapes[10],f'Actual changes: 12 → 36 reviewed entries; everyday questions, synonyms, reviewed negation and clearer unknown-claim guidance.\nExample: “Do antibiotics work for the flu?” is now matched.\nSeparate engineering benchmark: {benchmark["beforeCorrect"]}/20 → {benchmark["afterCorrect"]}/20 expected verdicts. This is not the participant retest.\n247 unit + 8 browser checks recorded.',10.8)
s.notes_slide.notes_text_frame.text=f'''The team explicitly supplied {summary['correct']} correct out of {summary['total']} claims from {summary['participants']} people. Arithmetic: {summary['accuracyPercent']:g}%; {summary['incorrect']} incorrect. Counts are team-reported; no individual tester identities, claim-level records, error wordings or human before/after outcomes were supplied. None have been invented. The separate automated benchmark actually ran 20 curated cases against baseline commit {benchmark['baselineCommit']} and current commit {benchmark['currentCommit']}: {benchmark['beforeCorrect']}/20 before, {benchmark['afterCorrect']}/20 after. It measures evidence coverage and wording support, not independent clinical accuracy or a five-person study. Old UNVERIFIED answers for missing coverage were appropriate abstentions. CSV/JSON reports are included in the local pack. Do not substitute the automated benchmark for human retest data.'''
reflections={
 'Nihit': 'Health Myth Buster helps people check health claims before sharing them. It matches claims to reviewed evidence and gives a verdict with source links. Unknown wording stays UNVERIFIED. I built and developed the tool and recorded an explanation of its purpose.',
 'Rudransh': 'Our tool helps families question unsupported health tips. It uses reviewed wording rules to show a verdict and evidence. I recorded the section about the common-cold claim returning FALSE. I explained why users should open the CDC and NHS source links.',
 'Priyushi': 'Our project checks general health claims using reviewed evidence. It says UNVERIFIED when a claim cannot be reliably matched. I recorded the unknown-claim example and explained that unknown is not automatically false. I also explained our transparent rules and separate pending-review candidates.',
 'Ahika': 'Our tool helps schools and NGOs teach health literacy. It gives explanations and sources, while the Testing Lab stores evaluations locally. I recorded the testing and closing section of our pitch. I explained that software checks are different from medical accuracy.'
}
s=prs.slides[8]
for name,idx in [('Nihit',7),('Rudransh',10),('Priyushi',13),('Ahika',16)]:
 fill(s.shapes[idx],reflections[name],10.2)
sh=s.shapes.add_textbox(Inches(.4),Inches(5.08),Inches(9.2),Inches(.16))
fill(sh,'Reflection drafts based on the recorded sections — each student must review and put the answer in their own words.',7.8,'9B5712')
s.notes_slide.notes_text_frame.text='These are clearly marked reflection drafts grounded in what each supplied recording actually says and in Nihit’s stated development role. They are not claimed to be independently written or approved by the students. Each student must understand, review and rewrite their own response before submitting the independent understanding check. No coding, research, testing or medical-review contributions have been invented for other members.'
file=OUT/'Health_Myth_Buster_Updated_Presentation.pptx';prs.save(file)
with ZipFile(BASE) as z:original_title=z.read('ppt/slides/slide1.xml')
with ZipFile(file) as z:parts=[(info,z.read(info.filename)) for info in z.infolist()]
tmp=file.with_suffix('.tmp')
with ZipFile(tmp,'w',ZIP_DEFLATED) as z:
 for info,data in parts:z.writestr(info,original_title if info.filename=='ppt/slides/slide1.xml' else data)
tmp.replace(file)
with ZipFile(file) as z:assert z.read('ppt/slides/slide1.xml')==original_title;assert z.testzip() is None
assert len(Presentation(file).slides)==9
# Reuse the existing preview renderer against the updated deck, without regenerating its content.
source=(ROOT/'scripts/presentation/build_deck.py').read_text()
render=source[source.index('def color(value'):source.index("(OUT/'deck-manifest.json')")]
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL
namespace={'base64':base64,'html':html,'prs':prs,'OUT':OUT,'INK':'1F2937','MSO_SHAPE_TYPE':MSO_SHAPE_TYPE,'MSO_FILL':MSO_FILL}
exec(compile(render,'deck_preview_renderer','exec'),namespace)
(OUT/'Reflections_For_Student_Review.md').write_text('# Reflection drafts for student review\n\nEach student must understand and rewrite their own answer. These drafts describe only the recorded presentation roles, not invented development contributions.\n\n'+'\n\n'.join('## '+name+'\n\n'+body for name,body in reflections.items()))
print('Updated PPT: 9 slides, title unchanged, team-reported results, real observed input and four fact-grounded reflection drafts.')
